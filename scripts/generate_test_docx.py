"""
Phase 0-A: 기존 YAML 정책 → .docx 역변환 테스트 데이터 생성기

기존 src/policies/*.yaml 파일에서 forbidden_words, compliance_checks, action 정보를 읽어
해당 정보가 명시된 한국식 내규 .docx 파일을 자동 생성한다.

사용법:
    python scripts/generate_test_docx.py
    python scripts/generate_test_docx.py --policy-dir src/policies --output-dir tests/fixtures/docx
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import yaml

sys.path.insert(0, str(Path(__file__).parent.parent))

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("ERROR: python-docx 패키지가 필요합니다. pip install python-docx")
    sys.exit(1)


# ── 정책 유형별 내규 템플릿 메타데이터 ──────────────────────────────────
_POLICY_META: dict[str, dict] = {
    "content":    {"dept": "준법감시팀", "category": "콘텐츠 안전"},
    "compliance": {"dept": "컴플라이언스팀", "category": "내부 준수"},
    "security":   {"dept": "정보보안팀", "category": "정보보안"},
    "privacy":    {"dept": "개인정보보호팀", "category": "개인정보"},
    "default":    {"dept": "법무팀", "category": "내부 규정"},
}

_ACTION_KO: dict[str, str] = {
    "BLOCK": "해당 내용을 즉시 차단하고 담당자에게 보고하여야 한다.",
    "LOG":   "해당 내용을 기록하고 모니터링 시스템에 이력을 남겨야 한다.",
}

_SEVERITY_KO: dict[str, str] = {
    "HIGH":   "높음 (즉각 대응 필요)",
    "MEDIUM": "중간 (정기 점검 대상)",
    "LOW":    "낮음 (참고 사항)",
}


def _guess_meta(policy_id: str) -> dict:
    pid_lower = policy_id.lower()
    for key, meta in _POLICY_META.items():
        if key in pid_lower:
            return meta
    return _POLICY_META["default"]


def _extract_forbidden_words(policy: dict) -> list[str]:
    words: list[str] = []
    for rule in policy.get("rules", []):
        cats = rule.get("parameters", {}).get("categories", {})
        for cat_data in cats.values():
            words.extend(cat_data.get("exact_terms", []))
    return list(dict.fromkeys(words))


def _extract_phrase_patterns(policy: dict) -> list[str]:
    patterns: list[str] = []
    for rule in policy.get("rules", []):
        cats = rule.get("parameters", {}).get("categories", {})
        for cat_data in cats.values():
            patterns.extend(cat_data.get("phrase_patterns", []))
    # regex 특수문자 제거해서 사람이 읽을 수 있는 형태로
    cleaned = []
    for p in patterns:
        readable = re.sub(r"\(\?i\)", "", p).strip()
        readable = re.sub(r"[\\^$.*+?()[\]{}|]", " ", readable).strip()
        if readable:
            cleaned.append(readable)
    return cleaned


def _extract_compliance_checks(policy: dict) -> list[dict]:
    """judge.criteria 텍스트에서 [CC-XXX] 항목 파싱."""
    criteria = policy.get("judge", {}).get("criteria", "") or ""
    checks = []
    for line in criteria.splitlines():
        m = re.search(r"\[CC-\d+\]\s*(.+?)(?:\s*\(심각도:\s*(\w+)\))?$", line.strip())
        if m:
            checks.append({
                "description": m.group(1).strip(),
                "severity": m.group(2) or "MEDIUM",
            })
    return checks


def _add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)


def _add_article(doc: Document, article_no: int, title: str, body: str) -> None:
    _add_heading(doc, f"제{article_no}조({title})", level=2)
    doc.add_paragraph(body)


def generate_docx(policy: dict, output_path: Path) -> None:
    policy_id   = policy.get("id", "UNKNOWN")
    policy_name = policy.get("name", "내부 규정")
    action_type = (policy.get("action") or {}).get("type", "LOG")
    meta        = _guess_meta(policy_id)

    forbidden_words  = _extract_forbidden_words(policy)
    phrase_patterns  = _extract_phrase_patterns(policy)
    compliance_items = _extract_compliance_checks(policy)

    doc = Document()

    # ── 제목 ──────────────────────────────────────────────────────
    title_para = doc.add_heading(f"{policy_name} 내부 규정", level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ── 메타 정보 ─────────────────────────────────────────────────
    meta_para = doc.add_paragraph()
    meta_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    meta_para.add_run(f"소관부서: {meta['dept']}  |  분류: {meta['category']}\n").italic = True
    meta_para.add_run(f"정책 ID: {policy_id}").italic = True
    doc.add_paragraph()

    # ── 제1조 목적 ────────────────────────────────────────────────
    _add_article(
        doc, 1, "목적",
        f"이 규정은 {policy_name}에 관한 필요한 사항을 정하여 회사의 건전한 "
        "업무 환경을 조성하고 임직원의 올바른 업무 수행을 지원함을 목적으로 한다.",
    )

    # ── 제2조 적용범위 ────────────────────────────────────────────
    _add_article(
        doc, 2, "적용범위",
        "이 규정은 회사의 모든 임직원, 계약직, 외주 용역 직원 및 "
        "회사 시스템을 이용하는 모든 사용자에게 적용된다.",
    )

    # ── 제3조 금지사항 ────────────────────────────────────────────
    _add_heading(doc, "제3조(금지사항)", level=2)

    if forbidden_words:
        fw_list = "、".join(f"「{w}」" for w in forbidden_words[:20])
        doc.add_paragraph(
            f"① 임직원은 회사 AI 시스템을 통해 다음 각 호의 금지 표현을 "
            f"사용하거나 생성하여서는 안 된다: {fw_list}."
        )

    if phrase_patterns:
        doc.add_paragraph(
            "② 다음의 금지 행위 패턴에 해당하는 내용은 사용을 엄격히 제한한다."
        )
        for i, pat in enumerate(phrase_patterns[:10], 1):
            doc.add_paragraph(f"{i}. {pat}", style="List Number")

    if not forbidden_words and not phrase_patterns:
        doc.add_paragraph(
            "임직원은 회사의 윤리 강령 및 관련 법규에 반하는 내용을 "
            "AI 시스템을 통해 생성하거나 전달하여서는 안 된다."
        )

    # ── 제4조 준수사항 ────────────────────────────────────────────
    if compliance_items:
        _add_heading(doc, "제4조(준수사항)", level=2)
        doc.add_paragraph(
            "임직원은 AI 시스템 사용 시 다음 각 호의 사항을 준수하여야 한다."
        )
        for i, item in enumerate(compliance_items[:15], 1):
            sev_ko = _SEVERITY_KO.get(item.get("severity", "MEDIUM"), "중간")
            doc.add_paragraph(
                f"{i}. {item['description']} [심각도: {sev_ko}]",
                style="List Number",
            )

    # ── 제5조 위반 시 처리 ────────────────────────────────────────
    article_no = 5 if compliance_items else 4
    _add_heading(doc, f"제{article_no}조(위반 시 처리)", level=2)
    doc.add_paragraph(
        f"① 이 규정을 위반한 경우 {_ACTION_KO.get(action_type, _ACTION_KO['LOG'])}"
    )
    doc.add_paragraph(
        "② 위반 사실이 확인된 경우 취업규칙 및 관련 법규에 따라 징계 조치할 수 있다."
    )

    # ── 부칙 ──────────────────────────────────────────────────────
    _add_heading(doc, "부칙", level=2)
    doc.add_paragraph("이 규정은 대표이사가 승인한 날로부터 시행한다.")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(f"  생성됨: {output_path.name}  (금지어 {len(forbidden_words)}개, 준수사항 {len(compliance_items)}개)")


def main() -> None:
    parser = argparse.ArgumentParser(description="YAML 정책 → .docx 역변환 생성기")
    parser.add_argument("--policy-dir", default="src/policies")
    parser.add_argument("--output-dir", default="tests/fixtures/docx")
    args = parser.parse_args()

    policy_dir = Path(args.policy_dir)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    yaml_files = list(policy_dir.glob("*.yaml"))
    if not yaml_files:
        print(f"YAML 파일 없음: {policy_dir}")
        sys.exit(1)

    print(f"\n[generate_test_docx] {len(yaml_files)}개 YAML → .docx 변환 시작\n")
    success = 0
    for yaml_path in sorted(yaml_files):
        try:
            with open(yaml_path, encoding="utf-8") as f:
                policy = yaml.safe_load(f) or {}
            if not isinstance(policy, dict):
                continue
            policy_id = policy.get("id") or yaml_path.stem
            out_path  = output_dir / f"{policy_id}_from_yaml.docx"
            generate_docx(policy, out_path)
            success += 1
        except Exception as e:
            print(f"  SKIP {yaml_path.name}: {e}")

    print(f"\n완료: {success}/{len(yaml_files)}개 생성 → {output_dir}\n")


if __name__ == "__main__":
    main()
